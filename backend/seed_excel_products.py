import os
import sys
import re
import json
import openpyxl

sys.path.append(os.path.dirname(os.path.abspath(__file__)))

from database import SessionLocal, engine, Base, get_mongo_db
import models

def format_hsn(raw):
    if not raw:
        return ""
    clean = re.sub(r'[^0-9]', '', str(raw).strip())
    if len(clean) == 8:
        return f"{clean[:4]}.{clean[4:6]}.{clean[6:]}"
    elif len(clean) == 6:
        return f"{clean[:4]}.{clean[4:6]}"
    elif len(clean) == 4:
        return clean
    return str(raw).strip()

def format_rate(val):
    if val is None or val == "":
        return None
    try:
        f = float(val)
        if f == 0:
            return "0%"
        elif f < 1:
            return f"{round(f * 100, 1)}%"
        else:
            return f"{int(f)}%" if f.is_integer() else f"{f}%"
    except Exception:
        return str(val).strip()

def main():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    
    # 1. Parse DUTY_2025 to build rate map
    excel_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'A3EXPRESS_10_26_27_DT_23_07_2026.xlsx')
    if not os.path.exists(excel_path):
        excel_path = os.path.join(os.getcwd(), 'A3EXPRESS_10_26_27_DT_23_07_2026.xlsx')
    
    duty_map = {}
    if os.path.exists(excel_path):
        print(f"Loading workbook for duty rates: {excel_path}")
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        if 'DUTY_2025' in wb.sheetnames:
            ws_duty = wb['DUTY_2025']
            duty_rows = list(ws_duty.iter_rows(values_only=True))
            for r in duty_rows[1:]:
                raw_hs = str(r[0]).strip() if r[0] is not None else ""
                desc = str(r[1]).strip() if len(r) > 1 and r[1] is not None else ""
                if not raw_hs or not re.sub(r'[^0-9]', '', raw_hs):
                    continue
                
                hsn_formatted = format_hsn(raw_hs)
                raw_digits = re.sub(r'[^0-9]', '', raw_hs)
                
                unit = str(r[13]).strip() if len(r) > 13 and r[13] is not None else "KG"
                gen_duty = format_rate(r[25]) if len(r) > 25 else None
                vat = format_rate(r[26]) if len(r) > 26 else None
                pal = format_rate(r[27]) if len(r) > 27 else None
                cess = format_rate(r[29]) if len(r) > 29 else None
                sscl = format_rate(r[32]) if len(r) > 32 else None
                scl = str(r[3]).strip() if len(r) > 3 and r[3] is not None else None
                
                info = {
                    "raw_hs": raw_digits,
                    "hs_code": hsn_formatted,
                    "description": desc,
                    "unit": unit,
                    "gen_duty": gen_duty or "20%",
                    "vat": vat or "18%",
                    "pal": pal or "Ex",
                    "cess": cess or "10%",
                    "sscl": sscl or "2.5%",
                    "scl": scl
                }
                duty_map[raw_digits] = info
                duty_map[hsn_formatted] = info
                if len(raw_digits) >= 4:
                    duty_map[raw_digits[:4]] = info

    # 2. Load 268 Master Products Catalog JSON
    master_json_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'brain', 'bdf3c8db-9e82-44e8-888e-4a11f38ce412', 'scratch', 'master_product_catalog.json')
    if not os.path.exists(master_json_path):
        # Fallback location
        master_json_path = r"C:\Users\mural\.gemini\antigravity-ide\brain\bdf3c8db-9e82-44e8-888e-4a11f38ce412\scratch\master_product_catalog.json"
    
    print(f"Loading master catalog from: {master_json_path}")
    with open(master_json_path, "r", encoding="utf-8") as f:
        master_products = json.load(f)

    print(f"Loaded {len(master_products)} products from master catalog.")

    # 3. Cache Chapters & Tariff Lines
    chapter_cache = {c.chapter_number: c for c in db.query(models.Chapter).all()}
    def get_or_create_chapter(chap_num):
        if chap_num in chapter_cache:
            return chapter_cache[chap_num]
        sec_num = 1
        if 1 <= chap_num <= 5: sec_num = 1
        elif 6 <= chap_num <= 14: sec_num = 2
        elif chap_num == 15: sec_num = 3
        elif 16 <= chap_num <= 24: sec_num = 4
        elif 25 <= chap_num <= 27: sec_num = 5
        elif 28 <= chap_num <= 38: sec_num = 6
        elif 39 <= chap_num <= 40: sec_num = 7
        else: sec_num = 8
        c = models.Chapter(
            section_number=sec_num,
            section_title=f"SECTION {sec_num}",
            chapter_number=chap_num,
            chapter_title=f"Customs Tariff Chapter {chap_num:02d}",
            source_pdf_filename=f"Tariff_Chap_{chap_num:02d}.pdf"
        )
        db.add(c)
        db.commit()
        db.refresh(c)
        chapter_cache[chap_num] = c
        return c

    tl_cache = {tl.hs_code: tl for tl in db.query(models.TariffLine).filter(models.TariffLine.hs_code != None).all()}
    existing_items = {ie.item_name.strip().lower(): ie for ie in db.query(models.ItemEntry).all()}

    inserted_count = 0
    updated_count = 0

    for prod in master_products:
        pname = prod['product_name'].strip()
        hsn_code = prod['hsn_code'].strip()
        category = prod['category'].strip()
        
        clean_name = pname.lower()
        digits = re.sub(r'[^0-9]', '', hsn_code)
        
        rates = duty_map.get(digits) or duty_map.get(hsn_code) or duty_map.get(digits[:4]) or {}
        gen_duty = rates.get('gen_duty', '20%')
        vat = rates.get('vat', '18%')
        pal = rates.get('pal', 'Ex')
        cess = rates.get('cess', '10%')
        sscl = rates.get('sscl', '2.5%')

        chap_num = int(digits[:2]) if len(digits) >= 2 else 1
        chap = get_or_create_chapter(chap_num)

        # Find or create TariffLine
        if hsn_code not in tl_cache:
            tl = models.TariffLine(
                chapter_id=chap.id,
                hs_code=hsn_code,
                description=pname,
                unit="KG" if "KG" in pname.upper() else "PCS",
                general_duty_rate=gen_duty,
                vat_rate=vat,
                pal_rate=pal,
                cess_rate=cess,
                sscl_rate=sscl,
                is_verified=True
            )
            db.add(tl)
            db.commit()
            db.refresh(tl)
            tl_cache[hsn_code] = tl
        else:
            tl = tl_cache[hsn_code]

        if clean_name in existing_items:
            ie = existing_items[clean_name]
            ie.hs_code = hsn_code
            ie.item_category = category
            ie.tariff_line_id = tl.id
            ie.tariff_description = pname
            ie.general_duty_rate = gen_duty
            ie.vat_rate = vat
            ie.pal_rate = pal
            ie.cess_rate = cess
            ie.sscl_rate = sscl
            ie.is_favorite = True
            updated_count += 1
        else:
            ie = models.ItemEntry(
                item_name=pname,
                item_category=category,
                item_classification="NORMAL",
                unit="KG" if "KG" in pname.upper() else "PCS",
                currency="LKR",
                tariff_line_id=tl.id,
                hs_code=hsn_code,
                tariff_description=pname,
                general_duty_rate=gen_duty,
                vat_rate=vat,
                pal_rate=pal,
                cess_rate=cess,
                sscl_rate=sscl,
                is_favorite=True
            )
            db.add(ie)
            inserted_count += 1

    db.commit()

    total_in_db = db.query(models.ItemEntry).count()
    print(f"\n==========================================")
    print(f"DATABASE SEEDING SUCCESSFUL!")
    print(f"Inserted: {inserted_count}")
    print(f"Updated: {updated_count}")
    print(f"Total Item Entries in Database: {total_in_db}")
    print(f"==========================================")

    # Sync to MongoDB Atlas if connected
    mongo_db = get_mongo_db()
    if mongo_db is not None:
        try:
            mongo_coll = mongo_db["item_entries"]
            print("Syncing master products to MongoDB Atlas...")
            for prod in master_products:
                pname = prod['product_name'].strip()
                hsn_code = prod['hsn_code'].strip()
                category = prod['category'].strip()
                digits = re.sub(r'[^0-9]', '', hsn_code)
                rates = duty_map.get(digits) or duty_map.get(hsn_code) or duty_map.get(digits[:4]) or {}
                
                mongo_coll.update_one(
                    {"item_name": pname},
                    {"$set": {
                        "item_name": pname,
                        "item_category": category,
                        "hs_code": hsn_code,
                        "unit": "KG" if "KG" in pname.upper() else "PCS",
                        "general_duty_rate": rates.get('gen_duty', '20%'),
                        "vat_rate": rates.get('vat', '18%'),
                        "pal_rate": rates.get('pal', 'Ex'),
                        "cess_rate": rates.get('cess', '10%'),
                        "sscl_rate": rates.get('sscl', '2.5%'),
                        "is_favorite": True
                    }},
                    upsert=True
                )
            print("MongoDB Atlas sync complete!")
        except Exception as e:
            print(f"MongoDB sync notice: {e}")

    db.close()

if __name__ == '__main__':
    main()
